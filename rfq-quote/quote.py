"""One synthetic text-PDF layout to an Excel quotation draft. See README.md."""
from pathlib import Path
from decimal import Decimal, InvalidOperation
from collections import defaultdict
import csv
import json
import re

from pypdf import PdfReader
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent


def extract_lines(path):
    text = "\n".join(page.extract_text(extraction_mode="layout") or "" for page in PdfReader(path).pages)
    rows, expected, ended = [], None, False
    for line in (value.strip() for value in text.splitlines() if value.strip()):
        if line in {"SYNTHETIC RFQ", "RFQ-DEMO-01"} or re.fullmatch(r"ITEM\s+CODE\s+DESCRIPTION\s+QTY\s+UNIT", line):
            continue
        if re.fullmatch(r"LINES \d+", line):
            expected = int(line.split()[1])
            continue
        if line == "END RFQ":
            ended = True
            continue
        match = re.fullmatch(r"(\d+)\s+(\S+)\s+(.+?)\s+(\S+)\s+(\S+)", line)
        if not match or ended:
            raise ValueError(f"Unrecognized sample layout line: {line!r}")
        number, code, description, quantity, unit = match.groups()
        rows.append(dict(line=int(number), code=code, description=description, quantity=quantity, unit=unit))
    if not ended or expected is None or len(rows) != expected or [r["line"] for r in rows] != list(range(1, expected + 1)):
        raise ValueError("Incomplete or non-sequential sample table")
    return rows, text


def read_csv(path):
    with Path(path).open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def prepare_quote(lines, catalogue, prices):
    products, approved = defaultdict(list), defaultdict(list)
    for product in catalogue:
        products[product["code"]].append(product)
    for price in prices:
        approved[price["code"]].append(price)
    versions = {p["price_list"] for p in prices}
    if len(versions) != 1 or not next(iter(versions), ""):
        raise ValueError("One named approved price list is required")
    version = next(iter(versions))
    result = []
    for line in lines:
        row = {**line, "status": "REVIEW", "reason": "", "unit_price_usd": "", "line_total_usd": "", "price_list": version}
        product, price = products[line["code"]], approved[line["code"]]
        if not re.fullmatch(r"[1-9][0-9]*", str(line["quantity"])):
            row["reason"] = "Quantity must be a positive whole number"
        elif len(product) != 1:
            row["reason"] = "Unknown product code" if not product else "Ambiguous catalogue code"
        elif product[0]["unit"] != line["unit"]:
            row["reason"] = "Requested unit differs from catalogue unit"
        elif len(price) != 1:
            row["reason"] = "No approved price" if not price else "Ambiguous approved price"
        elif price[0]["unit"] != line["unit"]:
            row["reason"] = "Price unit differs from requested unit"
        else:
            try:
                value = Decimal(price[0]["price_usd"])
                if not value.is_finite() or value < 0 or value.as_tuple().exponent < -2:
                    raise InvalidOperation
                amount = value * int(line["quantity"])
                row.update(status="MATCHED", reason="Exact code and unit; approved price", unit_price_usd=f"{value:.2f}", line_total_usd=f"{amount:.2f}")
            except InvalidOperation:
                row["reason"] = "Invalid approved USD price"
        result.append(row)
    subtotal = sum((Decimal(r["line_total_usd"]) for r in result if r["status"] == "MATCHED"), Decimal(0))
    reviews = sum(r["status"] == "REVIEW" for r in result)
    return {"state": "DRAFT FOR HUMAN APPROVAL", "price_list": version, "rows": result, "matched_lines": len(result) - reviews, "review_lines": reviews, "matched_subtotal_usd": f"{subtotal:.2f}", "quotation_total_usd": None if reviews else f"{subtotal:.2f}"}


def write_workbook(quote, output):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    for row in [["QUOTATION DRAFT", "SYNTHETIC DATA"], ["Status", quote["state"]], ["Price list", quote["price_list"]], ["Matched lines", quote["matched_lines"]], ["Lines needing review", quote["review_lines"]], ["Matched subtotal USD (incomplete)", quote["matched_subtotal_usd"]], ["Complete quotation total USD", quote["quotation_total_usd"] or "WITHHELD: unresolved lines"], ["Input", "RFQ-DEMO-01; one sample text-PDF layout"]]:
        summary.append(row)
    columns = ["line", "code", "description", "quantity", "unit", "status", "reason", "unit_price_usd", "line_total_usd", "price_list"]
    for name, selected in [("Draft", quote["rows"]), ("Review", [r for r in quote["rows"] if r["status"] == "REVIEW"])]:
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for row in selected:
            sheet.append([row[key] for key in columns])
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"
    for sheet in workbook:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="15392F")
        for cells in sheet.columns:
            letter = cells[0].column_letter
            sheet.column_dimensions[letter].width = min(56, max(15, max(len(str(c.value or "")) for c in cells) + 2))
        for row in sheet:
            for cell in row:
                # Source text remains literal text, including values beginning with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(output)


def run():
    lines, text = extract_lines(ROOT / "sample-rfq.pdf")
    quote = prepare_quote(lines, read_csv(ROOT / "catalogue.csv"), read_csv(ROOT / "approved-prices.csv"))
    write_workbook(quote, ROOT / "quotation-draft.xlsx")
    (ROOT / "extracted-text.txt").write_text(text, encoding="utf-8")
    (ROOT / "result.json").write_text(json.dumps(quote, indent=2) + "\n", encoding="utf-8")
    return quote


if __name__ == "__main__":
    print(json.dumps(run(), indent=2))
