"""Checks real PDF extraction, matching decisions and the written XLSX file."""
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from datetime import datetime, timezone
import json
import sys
import pypdf
import openpyxl
from quote import ROOT, extract_lines, read_csv, prepare_quote, write_workbook, run


def verify():
    checks = []
    def check(name, condition):
        assert condition, name
        checks.append({"name": name, "passed": True})

    result = run()
    check("Real PDF retains all eight requested lines", len(result["rows"]) == 8 and [r["line"] for r in result["rows"]] == list(range(1, 9)))
    check("Approved prices give three matches and USD66.75 incomplete subtotal", result["matched_lines"] == 3 and result["matched_subtotal_usd"] == "66.75")
    check("Unknown code, unit mismatch, missing price, ambiguous code and zero quantity each require review", [r["status"] for r in result["rows"]] == ["MATCHED", "MATCHED", "REVIEW", "REVIEW", "REVIEW", "REVIEW", "REVIEW", "MATCHED"] and len({r["reason"] for r in result["rows"] if r["status"] == "REVIEW"}) == 5)
    check("Unresolved quotation has no full total", result["quotation_total_usd"] is None)
    workbook = openpyxl.load_workbook(ROOT / "quotation-draft.xlsx")
    check("Written XLSX preserves all rows, five review entries and withheld total", workbook["Draft"].max_row == 9 and workbook["Review"].max_row == 6 and workbook["Summary"]["B7"].value == "WITHHELD: unresolved lines")

    lines, _ = extract_lines(ROOT / "sample-rfq.pdf")
    catalogue, prices = read_csv(ROOT / "catalogue.csv"), read_csv(ROOT / "approved-prices.csv")
    first = deepcopy(lines[:1])
    first[0]["price_usd"] = "0.01"
    check("An RFQ-supplied price cannot replace the approved list", prepare_quote(first, catalogue, prices)["matched_subtotal_usd"] == "28.50")
    duplicate = prices + [deepcopy(prices[0])]
    check("Two approved entries for the same code are flagged", prepare_quote(first, catalogue, duplicate)["rows"][0]["reason"] == "Ambiguous approved price")
    for bad in ["NaN", "Infinity", "-1", "1.001"]:
        changed = deepcopy(prices)
        changed[0]["price_usd"] = bad
        assert prepare_quote(first, catalogue, changed)["rows"][0]["status"] == "REVIEW"
    check("Non-finite, negative and over-precision prices require review", True)
    first[0]["quantity"] = "1.5"
    check("Fractional quantity is flagged under the declared whole-unit scope", prepare_quote(first, catalogue, prices)["rows"][0]["status"] == "REVIEW")
    first[0]["quantity"] = "1"
    first[0]["description"] = "=1+1"
    with TemporaryDirectory() as temporary:
        output = Path(temporary) / "literal.xlsx"
        write_workbook(prepare_quote(first, catalogue, prices), output)
        cell = openpyxl.load_workbook(output)["Draft"]["C2"]
        check("Source description is stored as literal text, not an Excel formula", cell.value == "=1+1" and cell.data_type == "s")

    if (ROOT / "second-rfq.pdf").exists():
        second, _ = extract_lines(ROOT / "second-rfq.pdf")
        quote = prepare_quote(second, catalogue, prices)
        check("A second PDF with changed quantities is extracted and priced independently", len(second) == 2 and quote["review_lines"] == 0 and quote["quotation_total_usd"] == "50.75")
    evidence = {"testedAt": datetime.now(timezone.utc).isoformat(), "platform": sys.platform, "python": sys.version.split()[0], "dependencies": {"pypdf": pypdf.__version__, "openpyxl": openpyxl.__version__}, "passed": len(checks), "failed": 0, "checks": checks, "scope": "Synthetic PDF/catalogue/Excel VERTICAL SLICE. No n8n runtime, client data or remote service involved."}
    (ROOT / "test-report.json").write_text(json.dumps(evidence, indent=2) + "\n", encoding="utf-8")
    return evidence


if __name__ == "__main__":
    print(json.dumps(verify(), indent=2))
